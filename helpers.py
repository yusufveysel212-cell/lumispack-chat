"""
LUMISPAC v3 - Yardimci Fonksiyonlar (helpers.py)

lumispac_app.py'den ayrilan saf yardimci fonksiyonlar.
Route degildir, is mantigi degildir.
"""
import io
import json
import hashlib
import re
import secrets
import uuid
import time
from datetime import datetime
from pathlib import Path

from flask import request, send_file, Response
from state import get_sessions, get_sessions_lock, _session_al, _session_bilgi
from utils import f, num, read_json, write_json, simdi_iso
from config import DATA_FILE, SET_FILE, QUOTE_FILE, ORDER_FILE, DEALER_FILE, DEFAULT_VOLUME
from data_repository import list_orders, list_quotes, list_dealers



# === normalize ===
def normalize(p):
    """Eksik alanları doldurur. Mevcut 0 değerleri DEĞİŞTİRMEZ."""
    sale = p.get('sale')
    volume = f(p.get('volume'), DEFAULT_VOLUME)
    volume_source = str(p.get('volume_source') or ('estimate' if volume > 0 else 'unknown')).lower()
    if volume_source not in ('unknown', 'estimate', 'measured'):
        volume_source = 'estimate' if volume > 0 else 'unknown'
    if volume <= 0:
        volume = 0.0
        volume_source = 'unknown'
    return {
        'id': str(p.get('id') or uuid.uuid4().hex[:12]),
        'code': str(p.get('code') or 'LUM-PAK'),
        'name': str(p.get('name') or '').strip(),
        'weight': f(p.get('weight')), 'qty': f(p.get('qty'), 1),
        'fire': f(p.get('fire')), 'shipping': f(p.get('shipping')), 'labor': f(p.get('labor')),
        'volume': volume, 'volume_source': volume_source,
        'sale': None if sale in (None, '') else f(sale),
        'stock': f(p.get('stock'), 0),
        'active': p.get('active', True) is not False,
    }



# === validate ===
def validate(items):
    if not isinstance(items, list):
        return 'Liste bekleniyor'
    for p in items:
        if not isinstance(p, dict):
            return 'Geçersiz kayıt'
        n = normalize(p)
        if not n['name']:
            return 'Ürün adı boş olamaz'
        if n['qty'] <= 0:
            return "%s: koli içi adet 0'dan büyük olmalı" % n['name']
        for k in ('weight', 'fire', 'shipping', 'labor', 'volume'):
            if n[k] < 0:
                return '%s: %s negatif olamaz' % (n['name'], k)
        if n['sale'] is not None and n['sale'] < 0:
            return '%s: satış fiyatı negatif olamaz' % n['name']
    return None



# === sig ===
def sig(p):
    n = str(p['name']).upper().replace('İ', 'I')
    m = re.search(r'\d+', n)
    flag = ('B' if 'BOMBEL' in n else '') + ('P' if 'PEYN' in n else '') + ('D' if 'DIKD' in n else '')
    return (int(m.group()) if m else 0), flag, round(f(p['weight']), 2), round(f(p['qty']))



# === same ===
def same(a, b):
    ca, fa, wa, qa = sig(a); cb, fb, wb, qb = sig(b)
    return fa == fb and wa == wb and qa == qb and abs(ca - cb) <= 1



# === merge_seed ===



# === startup_seed ===



# === calc ===
def calc(p, rate, raw_usd):
    gram = raw_usd * rate / 1000
    unit = p['weight'] * (1 + p['fire'] / 100) * gram
    koli = unit * p['qty'] + p['shipping'] + p['labor']
    return gram, unit, koli



#def parse_num(s):
    if s is None:
        return 0.0
    if not isinstance(s, str):
        s = str(s)
    s = s.strip().replace(' ', '')
    if not s:
        return 0.0
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0



# === norm_dealer ===
def norm_dealer(x):
    pr = x.get('prices') if isinstance(x.get('prices'), dict) else {}
    return {'id': str(x.get('id') or uuid.uuid4().hex[:10]), 'name': str(x.get('name') or '').strip()[:120],
            'contact': str(x.get('contact') or '')[:120], 'phone': str(x.get('phone') or '')[:40],
            'city': str(x.get('city') or '')[:60], 'address': str(x.get('address') or '')[:300],
            'discount': max(0.0, min(100.0, f(x.get('discount')))),
            'prices': {str(k): f(v) for k, v in pr.items() if f(v) > 0}}



# === order_lines ===
def order_lines(lines):
    out = []
    for l in lines[:200]:
        if isinstance(l, dict):
            out.append({'id': str(l.get('id', '')), 'code': str(l.get('code', ''))[:40], 'name': str(l.get('name', ''))[:120],
                        'weight': f(l.get('weight')), 'boxes': max(0, int(f(l.get('boxes')))), 'units': f(l.get('units')),
                        'price': f(l.get('price')), 'total': f(l.get('total')), 'cost': f(l.get('cost'))})
    return out



# === xlsx_response ===
def xlsx_response(headers, rows, name):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = 'Lumispac'
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='C8102E')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = '#,##0.00'
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(38, max(len(str(c.value or '')) for c in col) + 3)
    ws.freeze_panes = 'A2'
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')



# === _api_load_orders ===
def _api_load_orders():
    return list_orders()



# === _api_load_quotes ===
def _api_load_quotes():
    return list_quotes()



# === _api_load_dealers ===
def _api_load_dealers():
    return list_dealers()



# === _yeni_token ===
def _yeni_token():
    return secrets.token_urlsafe(32)



# === parse_num ===
def parse_num(s):
    """Guvenli sayi donusumu. None, '', gecersiz degerlerde 0.0 doner."""
    if s is None:
        return 0.0
    if not isinstance(s, str):
        s = str(s)
    s = s.strip().replace(' ', '')
    if not s:
        return 0.0
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0
